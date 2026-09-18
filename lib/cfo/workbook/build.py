"""C5 workbook builder: Excel files from a JSON spec."""
import datetime as _dt
import math
import os
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from cfo import xmlsafe
from cfo.brand import load_brand
from cfo.console import ToolkitError
from cfo.io import save_via_temp
from cfo.numbers import parse_number
from cfo.workbook import styles
from cfo.workbook.formula_policy import check_formula
from cfo.workbook.refs import SheetLayout, col_letter, resolve_formula, write_literal
from cfo.workbook.dashboard import write_dashboard
from cfo.workbook.rag import RAG_MESSAGE, status_formula, valid_rag

KINDS = ("table", "dashboard")
TYPES = ("text", "number", "percent", "date", "bool", "enum", "formula")
BAD_SHEET_CHARS = set("[]:*?/\\")
DATE_FORMAT = styles.DATE_FORMAT
MAX_WIDTH = 60
VALIDATION_SPARE_ROWS = 200
_KEY_RE = re.compile(r"^[a-z][a-z0-9_]*$")
_CELL_RE = re.compile(r"^\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}$")
# Excel limits (M8): data-validation error/prompt text and titles, cell text.
MAX_DV_TEXT = 255
MAX_DV_TITLE = 32
MAX_CELL_TEXT = 32767


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _check_columns_shape(sheet, where, problems):
    """I6: reject columns whose shape would otherwise crash the writer
    (openpyxl/comparisons), before any of that code runs."""
    columns = sheet.get("columns", [])
    if not isinstance(columns, list):
        problems.append((where, "columns must be a list"))
        return
    for i, col in enumerate(columns):
        cwhere = f"{where} columns[{i}]"
        if not isinstance(col, dict):
            problems.append((cwhere, "must be an object"))
            continue
        if "key" in col and col["key"] is not None and not isinstance(col["key"], str):
            problems.append((cwhere, "key must be a string"))
        if "label" in col and col["label"] is not None and not isinstance(col["label"], str):
            problems.append((cwhere, "label must be a string"))
        if "options" in col and col["options"] is not None and not (
                isinstance(col["options"], list) and
                all(isinstance(o, (str, int, float)) and not isinstance(o, bool) for o in col["options"])):
            problems.append((cwhere, "options must be a list of strings or numbers"))
        if "width" in col and col["width"] is not None and not (_is_number(col["width"]) and col["width"] > 0):
            problems.append((cwhere, "width must be a positive number"))
        if "format" in col and col["format"] is not None and not isinstance(col["format"], str):
            problems.append((cwhere, "format must be a string"))
        if "required" in col and col["required"] is not None and not isinstance(col["required"], bool):
            problems.append((cwhere, "required must be true or false"))


def _check_dashboard_shape(sheet, where, problems):
    """I6: reject dashboard blocks whose shape would otherwise crash the
    writer, before any of that code runs."""
    blocks = sheet.get("blocks", [])
    if not isinstance(blocks, list):
        problems.append((where, "blocks must be a list"))
        return
    for i, block in enumerate(blocks):
        bwhere = f"{where} blocks[{i}]"
        if not isinstance(block, dict):
            problems.append((bwhere, "must be an object"))
            continue
        for key in ("label", "text", "sort", "format", "source"):
            if key in block and block[key] is not None and not isinstance(block[key], str):
                problems.append((bwhere, f"{key} must be a string"))
        if "columns" in block and block["columns"] is not None and not (
                isinstance(block["columns"], list) and all(isinstance(c, str) for c in block["columns"])):
            problems.append((bwhere, "columns must be a list of strings"))


def _check_shapes(sheets, problems):
    """I6: validate the shape of every sheet, column and dashboard block the
    builder reads, before layout/writing code that assumes those shapes runs.
    `sheets` is already known to be a list of dicts with a valid name and
    kind (_check_sheet_names has raised otherwise)."""
    for sheet in sheets:
        where = sheet["name"]
        if "position" in sheet and not _is_number(sheet["position"]):
            # unlike the other optional fields, the sort below uses whatever
            # .get("position", ...) returns even when the key is present but
            # null, so null is not a safe "unset" here.
            problems.append((where, "position must be a number"))
        if sheet.get("kind", "table") == "table":
            if "freeze" in sheet and sheet["freeze"] is not None and not (
                    isinstance(sheet["freeze"], str) and _CELL_RE.match(sheet["freeze"])):
                problems.append((where, "freeze must be a cell reference such as B2"))
            if "rows" in sheet and not isinstance(sheet["rows"], list):
                problems.append((where, "rows must be a list"))
            _check_columns_shape(sheet, where, problems)
        else:
            _check_dashboard_shape(sheet, where, problems)


def _check_sheet_names(sheets, problems):
    seen = set()
    for i, sheet in enumerate(sheets):
        where = f"sheets[{i}]"
        name = sheet.get("name") if isinstance(sheet, dict) else None
        if not isinstance(name, str) or not name.strip():
            problems.append((where, "needs a name"))
            continue
        if len(name) > 31:
            problems.append((where, f"sheet name '{name}' is longer than 31 characters"))
        elif BAD_SHEET_CHARS & set(name):
            problems.append((where, f"sheet name '{name}' contains one of []:*?/\\"))
        elif name.startswith("'") or name.endswith("'"):
            problems.append((where, f"sheet name '{name}' must not start or end with '"))
        elif name.lower() in seen:
            problems.append((where, f"sheet name '{name}' is used twice"))
        elif name.startswith("_"):
            problems.append((where, "sheet names starting with _ are reserved"))
        elif sheet.get("kind", "table") not in KINDS:
            problems.append((where, f"kind must be one of {', '.join(KINDS)}"))
        seen.add(name.lower())


def _expand_columns(sheet, problems):
    columns, seen = [], set()
    for i, col in enumerate(sheet.get("columns", [])):
        where = f"{sheet['name']} columns[{i}]"
        key = col.get("key") if isinstance(col, dict) else None
        if not isinstance(key, str) or not _KEY_RE.match(key) or key in seen:
            problems.append((where, "needs a unique lower_case key"))
            continue
        ctype = col.get("type", "text")
        if ctype not in TYPES:
            problems.append((where, f"type must be one of {', '.join(TYPES)}"))
            continue
        if ctype == "enum" and not col.get("options"):
            problems.append((where, "enum columns need options"))
            continue
        if ctype == "formula":
            formula_text = str(col.get("formula", ""))
            if not formula_text.startswith("="):
                problems.append((where, "formula columns need a formula starting with ="))
                continue
            issue = check_formula(formula_text)
            if issue:
                problems.append((where, issue))
                continue
        rag = col.get("rag")
        if rag is not None and not valid_rag(rag):
            problems.append((where, RAG_MESSAGE))
            continue
        seen.add(key)
        label = col.get("label", key)
        columns.append(dict(col, type=ctype, label=label))
        if rag:
            seen.add(f"{key}_status")
            columns.append({"key": f"{key}_status", "label": f"{label} status", "type": "status",
                            "source": key, "rag": rag, "width": 14})
    if not columns:
        problems.append((sheet["name"], "needs at least one column"))
    return columns


def _convert(value, col):
    ctype = col["type"]
    if ctype == "text":
        return str(value), None
    if ctype in ("number", "percent"):
        text = value
        if ctype == "percent" and isinstance(value, str) and value.strip().endswith("%"):
            text = value.strip()[:-1]
        number = parse_number(text)
        if ctype == "percent":
            return number / 100, col.get("format", "0.00%")
        return number, col.get("format")
    if ctype == "date":
        try:
            return _dt.date.fromisoformat(str(value)[:10]), col.get("format", DATE_FORMAT)
        except ValueError:
            raise ValueError(f"'{value}' is not an ISO date (yyyy-mm-dd)")
    if ctype == "bool":
        if isinstance(value, bool):
            return value, None
        text = str(value).strip().lower()
        if text in ("true", "yes", "1"):
            return True, None
        if text in ("false", "no", "0"):
            return False, None
        raise ValueError(f"'{value}' is not true or false")
    if value not in col["options"]:
        raise ValueError(f"'{value}' is not one of {', '.join(str(o) for o in col['options'])}")
    return value, None


def _lists_sheet(wb):
    if "_lists" in wb.sheetnames:
        return wb["_lists"]
    ws = wb.create_sheet("_lists")
    ws.sheet_state = "hidden"
    return ws


def _shorten(text, limit):
    """Cut text to limit characters, returning (text, was_shortened)."""
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _validations(wb, ws, columns, layout, warnings):
    last = layout.last_row + VALIDATION_SPARE_ROWS
    for col in columns:
        letter = layout.columns[col["key"]]
        cells = f"{letter}{layout.first_row}:{letter}{last}"
        where = f"{layout.name}!{letter}1"
        if col["type"] == "enum":
            options = [str(o) for o in col["options"]]
            inline = ",".join(options)
            if len(inline) <= 255 and not any("," in o or '"' in o for o in options):
                dv = DataValidation(type="list", formula1=f'"{inline}"', allow_blank=True)
            else:
                lists = _lists_sheet(wb)
                column = 1 if lists["A1"].value is None else lists.max_column + 1
                for i, option in enumerate(options, 1):
                    write_literal(lists.cell(row=i, column=column), option)
                ref = col_letter(column)
                dv = DataValidation(type="list", formula1=f"'_lists'!${ref}$1:${ref}${len(options)}",
                                    allow_blank=True)
        elif col["type"] == "bool":
            dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
        else:
            continue
        error_title, title_cut = _shorten("Not in the list", MAX_DV_TITLE)
        error_text, text_cut = _shorten(f"Choose a value from the list for {col['label']}", MAX_DV_TEXT)
        dv.errorTitle = error_title
        dv.error = error_text
        if title_cut or text_cut:
            warnings.append((where, "the data-validation message was shortened to fit Excel's limit"))
        ws.add_data_validation(dv)
        dv.add(cells)


def _rag_rules(ws, columns, layout, brand):
    last = max(layout.last_row, layout.first_row)
    for col in columns:
        if col["type"] != "status":
            continue
        status = layout.columns[col["key"]]
        for letter in (layout.columns[col["source"]], status):
            cells = f"{letter}{layout.first_row}:{letter}{last}"
            for level in styles.RAG_LEVELS:
                ws.conditional_formatting.add(cells, FormulaRule(
                    formula=[f'${status}{layout.first_row}="{styles.RAG_LABELS[level]}"'],
                    fill=styles.rag_fill(brand, level), font=styles.rag_font(brand, level),
                    stopIfTrue=True))


def _widths(ws, columns, rows):
    for index, col in enumerate(columns, 1):
        width = col.get("width")
        if not width:
            values = [len(str(r.get(col["key"], ""))) for r in rows if isinstance(r, dict)]
            width = min(max([len(str(col["label"]))] + values) + 2, MAX_WIDTH)
        ws.column_dimensions[col_letter(index)].width = max(width, 8)


def _write_table(wb, ws, sheet, columns, layouts, brand, problems, warnings):
    name, layout, formulas = sheet["name"], layouts[sheet["name"]], 0
    for index, col in enumerate(columns, 1):
        cell = write_literal(ws.cell(row=1, column=index), col["label"])
        cell.font, cell.fill = styles.header_font(brand), styles.header_fill(brand)
        cell.alignment, cell.border = styles.header_alignment(), styles.BOX
    ws.row_dimensions[1].height = 30
    writable = {c["key"] for c in columns if c["type"] not in ("formula", "status")}
    rows = sheet.get("rows", [])
    for r, record in enumerate(rows, 2):
        if not isinstance(record, dict):
            problems.append((f"{name} row {r}", "must be an object"))
            continue
        for key in sorted(set(record) - writable):
            problems.append((f"{name} row {r}", f"unknown column '{key}'"))
        for index, col in enumerate(columns, 1):
            cell = ws.cell(row=r, column=index)
            cell.font, cell.border = styles.body_font(brand), styles.BOX
            where = f"{name}!{col_letter(index)}{r}"
            if col["type"] in ("formula", "status"):
                text = col["formula"] if col["type"] == "formula" else status_formula("{" + col["source"] + "}", col["rag"])
                try:
                    resolved = resolve_formula(text, name, r, layouts)
                except ValueError as exc:
                    problems.append((where, str(exc)))
                    continue
                issue = check_formula(resolved)
                if issue:
                    problems.append((where, issue))
                    continue
                cell.value = resolved
                formulas += 1
                if col.get("format"):
                    cell.number_format = col["format"]
                continue
            value = record.get(col["key"])
            if value is None or value == "":
                if col.get("required"):
                    warnings.append((where, f"{col['label']} is empty"))
                continue
            try:
                converted, number_format = _convert(value, col)
            except ValueError as exc:
                problems.append((where, str(exc)))
                continue
            if isinstance(converted, str) and len(converted) > MAX_CELL_TEXT:
                converted, _ = _shorten(converted, MAX_CELL_TEXT)
                warnings.append((where, f"text is longer than {MAX_CELL_TEXT:,} characters and was cut"))
            write_literal(cell, converted)
            if number_format:
                cell.number_format = number_format
    _validations(wb, ws, columns, layout, warnings)
    _rag_rules(ws, columns, layout, brand)
    _widths(ws, columns, rows)
    if sheet.get("freeze", "A2"):
        ws.freeze_panes = sheet.get("freeze", "A2")
    return formulas


def _save(wb, out_path):
    save_via_temp(out_path, wb.save, suffix=".xlsx")


def build_workbook(spec, out_path, brand=None):
    brand = brand or load_brand()
    spec, removed_chars = xmlsafe.clean_tree(spec)
    sheets = spec.get("sheets") if isinstance(spec, dict) else None
    if not isinstance(sheets, list) or not sheets:
        raise ToolkitError(("spec", "needs a non-empty 'sheets' list"))
    problems, warnings = [], []
    if removed_chars:
        warnings.append(("workbook", f"removed {removed_chars} character(s) that Excel cannot store"))
    _check_sheet_names(sheets, problems)
    if problems:
        raise ToolkitError(problems)
    _check_shapes(sheets, problems)
    if problems:
        raise ToolkitError(problems)
    layouts, columns = {}, {}
    for sheet in sheets:
        if sheet.get("kind", "table") == "table":
            columns[sheet["name"]] = _expand_columns(sheet, problems)
            layouts[sheet["name"]] = SheetLayout(
                sheet["name"], {c["key"]: col_letter(i) for i, c in enumerate(columns[sheet["name"]], 1)},
                2, 1 + len(sheet.get("rows", [])))
    if problems:
        raise ToolkitError(problems)

    wb = Workbook()
    wb.remove(wb.active)
    formulas = 0
    ordered = sorted(enumerate(sheets), key=lambda pair: (pair[1].get("position", 1000 + pair[0]), pair[0]))
    tables = {s["name"]: {"columns": columns[s["name"]], "rows": s.get("rows", [])}
              for s in sheets if s["name"] in columns}
    for _, sheet in ordered:
        ws = wb.create_sheet(sheet["name"])
        if sheet.get("kind", "table") == "dashboard":
            formulas += write_dashboard(ws, sheet, layouts, tables, brand, problems)
        else:
            formulas += _write_table(wb, ws, sheet, columns[sheet["name"]], layouts, brand,
                                     problems, warnings)
    if problems:
        raise ToolkitError(problems, warnings)
    _save(wb, out_path)
    return {"workbook": os.path.abspath(out_path), "formulas": formulas, "_warnings": warnings,
            "sheets": [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]}
