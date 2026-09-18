"""Dashboard sheets: title, KPI tiles, live tables and notes."""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from cfo.brand import hex6
from cfo.workbook import styles
from cfo.workbook.formula_policy import check_formula
from cfo.workbook.rag import RAG_MESSAGE, status_formula, valid_rag
from cfo.workbook.refs import col_letter, quote_sheet, resolve_formula, write_literal

BLOCK_TYPES = ("title", "kpi", "table", "note")
COLUMN_WIDTH = 22
MIN_COLUMN_WIDTH = 10
FORMULA_WIDTH = 12


def _anchor(at):
    letters, row = coordinate_from_string(str(at))
    return column_index_from_string(letters), row


def _kpi(ws, block, col, row, sheet, layouts, brand, problems, where):
    label = write_literal(ws.cell(row=row, column=col), block.get("label", ""))
    label.font = Font(name=brand["fonts"]["body"], size=10, color=hex6(brand["colours"]["muted"]))
    formula = str(block.get("formula", ""))
    if not formula.startswith("="):
        problems.append((where, "kpi needs a formula starting with ="))
        return 0
    try:
        resolved = resolve_formula(formula, sheet, None, layouts)
    except ValueError as exc:
        problems.append((where, str(exc)))
        return 0
    issue = check_formula(resolved)
    if issue:
        problems.append((where, issue))
        return 0
    value = ws.cell(row=row + 1, column=col, value=resolved)
    value.font = Font(name=brand["fonts"]["figures"], size=20, bold=True,
                      color=hex6(brand["colours"]["ink"]))
    if block.get("format"):
        value.number_format = block["format"]
    rules = block.get("rag")
    if not rules:
        return 1
    if not valid_rag(rules):
        problems.append((where, RAG_MESSAGE))
        return 1
    status = ws.cell(row=row + 1, column=col + 1, value=status_formula(value.coordinate, rules))
    for level in styles.RAG_LEVELS:
        ws.conditional_formatting.add(f"{value.coordinate}:{status.coordinate}", FormulaRule(
            formula=[f'${status.column_letter}${row + 1}="{styles.RAG_LABELS[level]}"'],
            fill=styles.rag_fill(brand, level), font=styles.rag_font(brand, level), stopIfTrue=True))
    return 2


def _column_width(cells):
    """Size to the content, not one width for every column. A formula's own text
    says nothing about how wide its result prints, so it counts as nominal."""
    lengths = [FORMULA_WIDTH if str(c.value).startswith("=") else len(str(c.value))
               for c in cells if c.value is not None]
    return min(max(max(lengths, default=0) + 4, MIN_COLUMN_WIDTH), COLUMN_WIDTH)


def _sort_value(value):
    if value is None or value == "":
        return (1, 0, "")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (0, 0, value)
    return (0, 1, str(value))


def _table(ws, block, col, row, layouts, tables, brand, problems, where):
    source = block.get("source")
    layout = layouts.get(source)
    if layout is None:
        problems.append((where, f"source '{source}' is not a table sheet"))
        return 0
    columns = {c["key"]: c for c in tables[source]["columns"]}
    keys = block.get("columns") or list(columns)
    unknown = [k for k in keys if k not in columns]
    sort_key = block.get("sort")
    if unknown or (sort_key and sort_key not in columns):
        problems.append((where, f"unknown column '{(unknown or [sort_key])[0]}' on sheet '{source}'"))
        return 0
    rows = tables[source]["rows"]
    order = list(range(len(rows)))
    if sort_key:
        order.sort(key=lambda i: _sort_value(rows[i].get(sort_key) if isinstance(rows[i], dict) else None))
    for j, key in enumerate(keys):
        cell = write_literal(ws.cell(row=row, column=col + j), columns[key]["label"])
        cell.font, cell.fill = styles.header_font(brand), styles.header_fill(brand)
        cell.alignment, cell.border = styles.header_alignment(), styles.BOX
    formulas = 0
    for i, source_index in enumerate(order, 1):
        for j, key in enumerate(keys):
            ref = f"{quote_sheet(source)}{layout.columns[key]}{layout.first_row + source_index}"
            cell = ws.cell(row=row + i, column=col + j, value=f'=IF({ref}="","",{ref})')
            cell.font, cell.border = styles.body_font(brand), styles.BOX
            spec = columns[key]
            if spec.get("format"):
                cell.number_format = spec["format"]
            elif spec["type"] == "date":
                cell.number_format = styles.DATE_FORMAT
            elif spec["type"] == "percent":
                cell.number_format = "0.00%"
            formulas += 1
    return formulas


def write_dashboard(ws, sheet, layouts, tables, brand, problems):
    ws.sheet_view.showGridLines = False
    formulas = 0
    for i, block in enumerate(sheet.get("blocks", [])):
        where = f"{sheet['name']} blocks[{i}]"
        btype = block.get("type") if isinstance(block, dict) else None
        if btype not in BLOCK_TYPES:
            problems.append((where, f"type must be one of {', '.join(BLOCK_TYPES)}"))
            continue
        try:
            col, row = _anchor(block.get("at", ""))
        except Exception:
            problems.append((where, "'at' must be a cell such as B4"))
            continue
        if btype == "title":
            cell = write_literal(ws.cell(row=row, column=col), block.get("text", ""))
            cell.font = Font(name=brand["fonts"]["heading"], size=18, bold=True,
                             color=hex6(brand["colours"]["ink"]))
        elif btype == "note":
            cell = write_literal(ws.cell(row=row, column=col), block.get("text", ""))
            cell.font = Font(name=brand["fonts"]["body"], size=9, italic=True,
                             color=hex6(brand["colours"]["muted"]))
        elif btype == "kpi":
            formulas += _kpi(ws, block, col, row, sheet["name"], layouts, brand, problems, where)
        else:
            formulas += _table(ws, block, col, row, layouts, tables, brand, problems, where)
    for index in range(1, ws.max_column + 1):
        letter = col_letter(index)
        ws.column_dimensions[letter].width = _column_width(ws[letter])
    return formulas
