"""C6 Word builder: .docx from Markdown, in Hedj branding."""
import math
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from cfo import xmlsafe
from cfo.brand import hex6, load_brand
from cfo.console import ToolkitError
from cfo.documents.markdown import parse_front_matter, parse_md, runs
from cfo.io import save_via_temp

TEMPLATES = {
    "report": {"margins_mm": (30, 25, 22, 20), "body_pt": 10.5, "chars_per_line": 87, "lines_per_page": 46},
    "policy": {"margins_mm": (30, 25, 22, 20), "body_pt": 10.5, "chars_per_line": 87, "lines_per_page": 46},
    "onepager": {"margins_mm": (18, 15, 15, 14), "body_pt": 9.5, "chars_per_line": 112, "lines_per_page": 60},
}
# Hairline between table rows: the same grey the workbook builder uses.
HAIRLINE = "D1D5DB"
# A figure once currency, thousands separators, a sign, brackets or % are stripped.
NUMERIC_RE = re.compile(r"^[-+(]?\s*[€£$]?\s*\d[\d,\s]*(?:\.\d+)?\s*%?\s*\)?$")
LIST_STYLES = {(True, 0): "List Number", (True, 1): "List Number 2",
               (False, 0): "List Bullet", (False, 1): "List Bullet 2"}
SAFE_LINK_SCHEMES = ("http://", "https://", "mailto:")


def _rgb(colour):
    return RGBColor.from_string(hex6(colour))


def _font(style, name, size=None, colour=None, bold=None):
    style.font.name = name
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        style.font.size = Pt(size)
    if colour:
        style.font.color.rgb = _rgb(colour)
    if bold is not None:
        style.font.bold = bold


def _setup(doc, brand, settings):
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    left, right, top, bottom = settings["margins_mm"]
    section.left_margin, section.right_margin = Mm(left), Mm(right)
    section.top_margin, section.bottom_margin = Mm(top), Mm(bottom)
    colours, fonts = brand["colours"], brand["fonts"]
    _font(doc.styles["Normal"], fonts["body"], settings["body_pt"], colours["ink"])
    body = doc.styles["Normal"].paragraph_format
    body.space_after, body.widow_control = Pt(6), True
    _font(doc.styles["Title"], fonts["heading"], 24, colours["ink"], True)
    doc.styles["Title"].paragraph_format.space_after = Pt(2)
    _font(doc.styles["Subtitle"], fonts["heading"], 13, colours["muted"])
    for level, size, before, after in ((1, 16, 18, 6), (2, 13, 14, 4), (3, 11.5, 12, 3)):
        _font(doc.styles[f"Heading {level}"], fonts["heading"], size,
              colours["primary"] if level == 1 else colours["ink"], True)
        fmt = doc.styles[f"Heading {level}"].paragraph_format
        fmt.space_before, fmt.space_after, fmt.keep_with_next = Pt(before), Pt(after), True


def _field(paragraph, instruction):
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"}, f" {instruction} "),
                             ("w:fldChar", {"w:fldCharType": "separate"}, None),
                             ("w:t", {}, "1"),
                             ("w:fldChar", {"w:fldCharType": "end"}, None)):
        el = OxmlElement(tag)
        for key, value in attrs.items():
            el.set(qn(key), value)
        if text is not None:
            el.text = text
        paragraph.add_run()._r.append(el)


def _header_footer(doc, meta, brand):
    section = doc.sections[0]
    muted = _rgb(brand["colours"]["muted"])
    header = section.header.paragraphs[0]
    header.text = " · ".join(str(v) for v in (meta.get("title"), meta.get("classification")) if v)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size, run.font.color.rgb = Pt(8), muted
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")
    footer.add_run(f"    {brand['footer_text']}")
    for run in footer.runs:
        run.font.size, run.font.color.rgb = Pt(8), muted
    if meta.get("disclaimer", True) and brand.get("disclaimer"):
        note = section.footer.add_paragraph(brand["disclaimer"])
        for run in note.runs:
            run.font.size, run.font.color.rgb = Pt(7), muted


def _title_block(doc, meta, brand):
    if meta.get("title"):
        doc.add_paragraph(str(meta["title"]), style="Title")
    if meta.get("subtitle"):
        doc.add_paragraph(str(meta["subtitle"]), style="Subtitle")
    line = " · ".join(str(meta[k]) for k in ("company", "date") if meta.get(k))
    if line:
        run = doc.add_paragraph().add_run(line)
        run.font.color.rgb = _rgb(brand["colours"]["muted"])


def _hyperlink(paragraph, text, url, brand):
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
    colour, underline = OxmlElement("w:color"), OxmlElement("w:u")
    colour.set(qn("w:val"), hex6(brand["colours"]["primary"]))
    underline.set(qn("w:val"), "single")
    rpr.append(colour)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _add_runs(paragraph, text, brand, warnings=None):
    for part, bold, italic, code, link in runs(text):
        if link:
            if str(link).lower().startswith(SAFE_LINK_SCHEMES):
                _hyperlink(paragraph, part, link, brand)
            else:
                run = paragraph.add_run(part)
                run.bold = bold or None
                run.italic = italic or None
                if warnings is not None:
                    warnings.append((link, "link removed: only http, https and mailto links are kept"))
            continue
        run = paragraph.add_run(part)
        run.bold = bold or None
        run.italic = italic or None
        if code:
            run.font.name = "Consolas"


def _shade(cell, colour):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6(colour))
    cell._tc.get_or_add_tcPr().append(shd)


def _cell_borders(cell, colour, size):
    """Rule under the cell only: no vertical lines, no grid."""
    borders = OxmlElement("w:tcBorders")
    edges = (("w:top", {"w:val": "nil"}), ("w:left", {"w:val": "nil"}), ("w:right", {"w:val": "nil"}),
             ("w:bottom", {"w:val": "single", "w:sz": size, "w:space": "0", "w:color": colour}))
    for tag, attrs in edges:
        el = OxmlElement(tag)
        for key, value in attrs.items():
            el.set(qn(key), value)
        borders.append(el)
    cell._tc.get_or_add_tcPr().append(borders)


def _cell_padding(table, dxa="72"):
    """Without the grid, text needs room above and below each rule."""
    margins = OxmlElement("w:tblCellMar")
    for tag in ("w:top", "w:bottom"):
        el = OxmlElement(tag)
        el.set(qn("w:w"), dxa)
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    table._tbl.tblPr.append(margins)


def _numeric_columns(rows, cols):
    """A column is numeric when every non-empty body cell reads as a figure."""
    numeric = []
    for c in range(cols):
        values = [row[c].strip() for row in rows[1:] if c < len(row) and row[c].strip()]
        numeric.append(bool(values) and all(NUMERIC_RE.match(v) for v in values))
    return numeric


def _table(doc, rows, brand, usable_mm, warnings=None):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_padding(table)
    numeric = _numeric_columns(rows, cols)
    lengths = [max(max((len(r[c]) if c < len(r) else 0) for r in rows), 4) for c in range(cols)]
    widths = [Mm(usable_mm * n / sum(lengths)) for n in lengths]
    for r, row in enumerate(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.width = widths[c]
            paragraph = cell.paragraphs[0]
            _add_runs(paragraph, row[c] if c < len(row) else "", brand, warnings)
            if numeric[c]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if r == 0:
                _shade(cell, brand["colours"]["primary"])
                _cell_borders(cell, hex6(brand["colours"]["ink"]), "8")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = _rgb(brand["colours"]["white"])
            else:
                _cell_borders(cell, HAIRLINE, "4")
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)


def _rule(doc, brand):
    paragraph = doc.add_paragraph()
    borders, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"),
                       ("w:color", hex6(brand["colours"]["muted"]))):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _callout(doc, text, brand, warnings=None):
    """A left-ruled, tinted note: Word's Intense Quote reads as a Word template."""
    paragraph = doc.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    borders, left = OxmlElement("w:pBdr"), OxmlElement("w:left")
    for key, value in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"),
                       ("w:color", hex6(brand["colours"]["primary"]))):
        left.set(qn(key), value)
    borders.append(left)
    ppr.append(borders)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6(brand["colours"]["light"]))
    ppr.append(shd)
    fmt = paragraph.paragraph_format
    fmt.left_indent, fmt.space_before = Mm(4), Pt(8)
    _add_runs(paragraph, text, brand, warnings)
    return paragraph


POLICY_FIELDS = (("version", "Version"), ("owner", "Policy owner"), ("approved_by", "Approved by"),
                 ("approved_on", "Approval date"), ("next_review", "Next review"))


def _policy_front(doc, meta, brand, warnings=None):
    rows = [["Document control", ""]] + [[label, str(meta.get(key, "To be confirmed"))]
                                         for key, label in POLICY_FIELDS]
    _table(doc, rows, brand, 120, warnings)
    doc.add_paragraph()


def _number_headings(doc, brand):
    """Number Heading 1-3 as 1, 1.1, 1.1.1 with a multi-level list linked to the styles."""
    numbering = doc.part.numbering_part.element
    ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for ilvl, text in enumerate(("%1", "%1.%2", "%1.%2.%3")):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:pStyle", f"Heading{ilvl + 1}"),
                           ("w:lvlText", text), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), value)
            lvl.append(el)
        ppr, ind = OxmlElement("w:pPr"), OxmlElement("w:ind")
        ind.set(qn("w:left"), str(567 * (ilvl + 1)))
        ind.set(qn("w:hanging"), "567")
        ppr.append(ind)
        lvl.append(ppr)
        # w:pPr must precede w:rPr inside w:lvl, or Word offers to repair the file.
        rpr, fonts = OxmlElement("w:rPr"), OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), brand["fonts"]["figures"])
        rpr.append(fonts)
        lvl.append(rpr)
        abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)      # abstractNum elements must precede num elements
    else:
        numbering.append(abstract)
    num = numbering.add_num(abstract_id)
    for level in (1, 2, 3):
        numpr = doc.styles[f"Heading {level}"].element.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = level - 1
        numpr.get_or_add_numId().val = num.numId
    return num.numId


def _new_list_num(doc, style_name):
    """A fresh numbering instance for a numbered list, restarting at 1."""
    ppr = doc.styles[style_name].element.pPr
    numpr = ppr.numPr if ppr is not None else None
    if numpr is None or numpr.numId is None:
        return None
    numbering = doc.part.numbering_part.element
    abstract_id = numbering.num_having_numId(numpr.numId.val).abstractNumId.val
    num = numbering.add_num(abstract_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId


def _list_item(doc, block, state, brand, warnings=None):
    _, text, ordered, depth = block
    style = LIST_STYLES[(ordered, depth)]
    paragraph = doc.add_paragraph(style=style)
    _add_runs(paragraph, text, brand, warnings)
    for deeper in [k for k in state if k[1] > depth]:
        del state[deeper]          # a nested list restarts each time it reappears
    if not ordered:
        return
    key = (style, depth)
    if key not in state:
        state[key] = _new_list_num(doc, style)
    if state[key] is not None:
        numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = state[key]


def estimate_pages(blocks, settings, meta):
    per_line, per_page = settings["chars_per_line"], settings["lines_per_page"]
    lines = 4 if meta.get("title") else 0
    for block in blocks:
        kind = block[0]
        if kind == "h":
            lines += 2
        elif kind in ("p", "quote"):
            lines += math.ceil(len(block[1]) / per_line) + 1
        elif kind == "li":
            lines += math.ceil(len(block[1]) / (per_line - 6))
        elif kind == "table":
            lines += sum(max(1, math.ceil(sum(len(c) for c in row) / per_line)) for row in block[1]) + 1
        elif kind == "image":
            lines += 15
        elif kind == "hr":
            lines += 1
        elif kind == "pagebreak":
            lines = math.ceil(lines / per_page) * per_page
    return max(1, math.ceil(lines / per_page))


def _save(doc, out_path):
    save_via_temp(out_path, doc.save, suffix=".docx")


def build_docx(md_text, out_path, template=None, base_dir=None, brand=None):
    brand = brand or load_brand()
    md_text, removed_chars = xmlsafe.clean(md_text)
    meta, body = parse_front_matter(md_text)
    template = template or meta.get("template") or "report"
    if template not in TEMPLATES:
        raise ToolkitError(("template", f"must be one of {', '.join(TEMPLATES)}"))
    settings = TEMPLATES[template]
    blocks = parse_md(body)
    doc = Document()
    warnings = []
    if removed_chars:
        warnings.append(("document", f"removed {removed_chars} character(s) that Word cannot store"))
    _setup(doc, brand, settings)
    _header_footer(doc, meta, brand)
    _title_block(doc, meta, brand)
    if template == "policy":
        _policy_front(doc, meta, brand, warnings)
        _number_headings(doc, brand)
    usable_mm = 210 - settings["margins_mm"][0] - settings["margins_mm"][1]
    root = os.path.realpath(base_dir or os.getcwd())
    list_state = {}
    for block in blocks:
        kind = block[0]
        if kind != "li":
            list_state.clear()
        if kind == "h":
            _add_runs(doc.add_heading(level=min(block[1], 4)), block[2], brand, warnings)
        elif kind == "p":
            _add_runs(doc.add_paragraph(), block[1], brand, warnings)
        elif kind == "li":
            _list_item(doc, block, list_state, brand, warnings)
        elif kind == "quote":
            _callout(doc, block[1], brand, warnings)
        elif kind == "table":
            _table(doc, block[1], brand, usable_mm, warnings)
        elif kind == "image":
            path = os.path.realpath(os.path.join(root, block[2]))
            try:
                inside = os.path.commonpath([os.path.normcase(root), os.path.normcase(path)]) == os.path.normcase(root)
            except ValueError:
                inside = False
            if not inside:
                warnings.append((block[2], "image outside the document folder, left out"))
            elif not os.path.isfile(path):
                warnings.append((block[2], "image not found, left out"))
            else:
                try:
                    doc.add_picture(path, width=Mm(usable_mm))
                except Exception:  # noqa: BLE001 - any bad/corrupt image must warn, not crash (I6)
                    warnings.append((block[2], "image could not be read, left out"))
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "hr":
            _rule(doc, brand)
    pages = estimate_pages(blocks, settings, meta)
    if template == "onepager" and pages > 1:
        warnings.append(("document", f"estimated {pages} pages; a one-pager should fit on one page"))
    _save(doc, out_path)
    return {"document": os.path.abspath(out_path), "template": template, "pages_estimate": pages,
            "_warnings": warnings}
