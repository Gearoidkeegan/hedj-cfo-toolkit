"""C6 deck builder: 16:9 PowerPoint decks from a JSON spec, in Hedj branding."""
import math
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from cfo import xmlsafe
from cfo.brand import hex6, load_brand
from cfo.console import ToolkitError
from cfo.documents.fit import fit_items
from cfo.io import save_via_temp

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MAX_SLIDES = 12
LAYOUTS = ("title", "messages", "text_table", "metrics", "chart")
CHART_TYPES = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
               "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}
SERIES_COLOURS = ("#1DB88D", "#1F2A37", "#FFC000", "#7C6AED", "#E5484D")
HAIRLINE = "D1D5DB"


def _is_text(value):
    return isinstance(value, str) and value.strip() != ""


def _is_finite_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def validate_slide(index, slide):
    where = f"slides[{index}]"
    if not isinstance(slide, dict) or slide.get("layout") not in LAYOUTS:
        return [(where, f"layout must be one of {', '.join(LAYOUTS)}")]
    layout = slide["layout"]
    if "footer" in slide and slide["footer"] is not None and not isinstance(slide["footer"], str):
        return [(where, "footer must be a string")]
    if layout == "title":
        return [] if _is_text(slide.get("title")) else [(where, "title slides need a title")]
    if not _is_text(slide.get("headline")):
        return [(where, "needs a headline")]
    if layout == "messages":
        items = slide.get("messages")
        if not isinstance(items, list) or not 3 <= len(items) <= 5 or not all(_is_text(m) for m in items):
            return [(where, "messages needs 3 to 5 items")]
    elif layout == "text_table":
        bullets, table = slide.get("bullets", []), slide.get("table")
        if not isinstance(bullets, list) or len(bullets) > 6:
            return [(where, "bullets allows at most 6 items")]
        if not isinstance(table, list) or not 2 <= len(table) <= 9 or \
                not all(isinstance(r, list) and 1 <= len(r) <= 5 for r in table):
            return [(where, "table needs a header row plus 1 to 8 rows of at most 5 columns")]
    elif layout == "metrics":
        tiles = slide.get("tiles")
        if not isinstance(tiles, list) or not 2 <= len(tiles) <= 6 or not all(
                isinstance(t, dict) and _is_text(t.get("label")) and t.get("value") not in (None, "")
                for t in tiles):
            return [(where, "tiles needs 2 to 6 items, each with a label and a value")]
        bad = [t["rag"] for t in tiles if t.get("rag") not in (None, "red", "amber", "green")]
        if bad:
            return [(where, f"tile rag '{bad[0]}' must be red, amber or green")]
    elif layout == "chart":
        chart = slide.get("chart")
        chart_type = chart.get("type") if isinstance(chart, dict) else None
        if not isinstance(chart, dict) or not isinstance(chart_type, str) or chart_type not in CHART_TYPES:
            return [(where, f"chart type must be one of {', '.join(CHART_TYPES)}")]
        categories, series = chart.get("categories"), chart.get("series")
        if not isinstance(categories, list) or not categories or not all(
                isinstance(c, (str, int, float)) and not isinstance(c, bool) for c in categories):
            return [(where, "chart needs categories, each a string or a number")]
        if not isinstance(series, list) or not series:
            return [(where, "chart needs at least one series")]
        if chart["type"] == "pie" and len(series) != 1:
            return [(where, "a pie chart takes exactly one series")]
        for s in series:
            values = s.get("values") if isinstance(s, dict) else None
            name = s.get("name", "?") if isinstance(s, dict) else "?"
            if not isinstance(values, list) or len(values) != len(categories):
                return [(where, f"series '{name}' needs one number per category ({len(categories)})")]
            if not all(_is_finite_number(v) for v in values):
                return [(where, f"series '{name}' values must be finite numbers")]
    return []


def _rgb(colour):
    return RGBColor.from_string(hex6(colour))


def _text(slide, box, text, size, colour, brand, bold=False, align=PP_ALIGN.LEFT, heading=False,
          figures=False):
    left, top, width, height = (Inches(v) for v in box)
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.size, run.font.bold = Pt(size), bold
    run.font.color.rgb = _rgb(colour)
    run.font.name = brand["fonts"]["figures" if figures else ("heading" if heading else "body")]
    return shape


def _rect(slide, box, colour, shape_type=MSO_SHAPE.RECTANGLE):
    left, top, width, height = (Inches(v) for v in box)
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(colour)
    shape.line.fill.background()
    return shape


def _fit(items, box, size, warnings, where):
    fitted, pt, _ = fit_items(items, box, size=size)
    for name, before, after in zip(where, items, fitted):
        if " ".join(str(before).split()) != after:
            warnings.append((name, "shortened to fit"))
    return fitted, pt


def _chrome(slide, number, headline, footer, brand, warnings):
    colours = brand["colours"]
    _rect(slide, (0, 0, 13.333, 0.12), colours["primary"])
    if headline is not None:
        (text,), pt = _fit([headline], (12.1, 0.9), 26, warnings, [f"slide {number} headline"])
        _text(slide, (0.6, 0.35, 12.1, 0.9), text, pt, colours["ink"], brand, bold=True, heading=True)
    _text(slide, (0.6, 7.0, 10.8, 0.35), footer, 9, colours["muted"], brand)
    _text(slide, (12.0, 7.0, 0.8, 0.35), str(number), 9, colours["muted"], brand, align=PP_ALIGN.RIGHT)


def _title_slide(slide, spec, number, brand, warnings):
    colours = brand["colours"]
    _rect(slide, (0.6, 2.4, 0.12, 2.2), colours["accent"])
    (title,), pt = _fit([spec["title"]], (11.5, 1.6), 40, warnings, [f"slide {number} title"])
    _text(slide, (0.95, 2.3, 11.5, 1.6), title, pt, colours["ink"], brand, bold=True, heading=True)
    if spec.get("subtitle"):
        _text(slide, (0.95, 3.9, 11.5, 0.6), spec["subtitle"], 20, colours["muted"], brand)
    if spec.get("date"):
        _text(slide, (0.95, 4.5, 11.5, 0.5), spec["date"], 14, colours["muted"], brand)


def _messages_slide(slide, spec, number, brand, warnings):
    colours = brand["colours"]
    items = spec["messages"]
    names = [f"slide {number} messages[{i}]" for i in range(len(items))]
    fitted, pt = _fit(items, (11.0, 5.2), 20, warnings, names)
    row = 5.2 / len(items)
    for i, message in enumerate(fitted):
        top = 1.5 + i * row
        circle = _rect(slide, (0.6, top + 0.05, 0.55, 0.55), colours["primary"], MSO_SHAPE.OVAL)
        circle.text_frame.text = str(i + 1)
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        circle.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        circle.text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb(colours["white"])
        _text(slide, (1.4, top, 11.0, row - 0.1), message, pt, colours["ink"], brand)


def _text_table_slide(slide, spec, number, brand, warnings):
    colours = brand["colours"]
    bullets = spec.get("bullets", [])
    if bullets:
        names = [f"slide {number} bullets[{i}]" for i in range(len(bullets))]
        fitted, pt = _fit(bullets, (5.6, 5.2), 18, warnings, names)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(5.6), Inches(5.2))
        frame = box.text_frame
        frame.word_wrap = True
        for i, bullet in enumerate(fitted):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = f"•  {bullet}"
            run.font.size, run.font.color.rgb = Pt(pt), _rgb(colours["ink"])
            run.font.name = brand["fonts"]["body"]
            paragraph.space_after = Pt(8)
    rows = spec["table"]
    cols = max(len(r) for r in rows)
    shape = slide.shapes.add_table(len(rows), cols, Inches(6.6), Inches(1.5), Inches(6.1),
                                   Inches(0.45 * len(rows)))
    table = shape.table
    table.first_row, table.horz_banding = False, False
    for r, row in enumerate(rows):
        for c in range(cols):
            value = str(row[c]) if c < len(row) else ""
            if len(value) > 60:
                warnings.append((f"slide {number} table[{r}][{c}]", "shortened to fit"))
                value = value[:59] + "…"
            cell = table.cell(r, c)
            cell.text = value
            run = cell.text_frame.paragraphs[0].runs[0] if value else None
            if run is not None:
                run.font.size = Pt(12)
                run.font.name = brand["fonts"]["body"]
                run.font.bold = r == 0
                run.font.color.rgb = _rgb(colours["white"] if r == 0 else colours["ink"])
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(colours["primary"] if r == 0 else colours["white"])


def _metrics_slide(slide, spec, number, brand, warnings):
    colours = brand["colours"]
    tiles = spec["tiles"]
    per_row = 2 if len(tiles) in (2, 4) else 3
    width = (12.1 - 0.3 * (per_row - 1)) / per_row
    for i, tile in enumerate(tiles):
        row, col = divmod(i, per_row)
        left, top = 0.6 + col * (width + 0.3), 1.6 + row * 2.6
        _rect(slide, (left, top, width, 2.3), colours["light"], MSO_SHAPE.ROUNDED_RECTANGLE)
        _rect(slide, (left, top, 0.12, 2.3), colours[tile.get("rag") or "primary"])
        _text(slide, (left + 0.35, top + 0.2, width - 0.5, 0.5), str(tile["label"]), 14, colours["muted"], brand)
        _text(slide, (left + 0.35, top + 0.7, width - 0.5, 0.9), str(tile["value"]), 32, colours["ink"],
              brand, bold=True, figures=True)
        if tile.get("delta"):
            _text(slide, (left + 0.35, top + 1.6, width - 0.5, 0.5), str(tile["delta"]), 12,
                  colours["muted"], brand)


def _chart_slide(slide, spec, number, brand, warnings):
    colours = brand["colours"]
    chart_spec = spec["chart"]
    data = CategoryChartData()
    data.categories = [str(c) for c in chart_spec["categories"]]
    for series in chart_spec["series"]:
        data.add_series(str(series.get("name", "")), series["values"])
    width = 8.4 if spec.get("takeaway") else 12.1
    frame = slide.shapes.add_chart(CHART_TYPES[chart_spec["type"]], Inches(0.6), Inches(1.4),
                                   Inches(width), Inches(5.3), data)
    chart = frame.chart
    chart.has_legend = len(chart_spec["series"]) > 1 or chart_spec["type"] == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.font.name = brand["fonts"]["body"]
    chart.font.size = Pt(11)
    if chart_spec["type"] != "pie":
        value_axis = chart.value_axis
        value_axis.has_major_gridlines = True
        gridline = value_axis.major_gridlines.format.line
        gridline.color.rgb = RGBColor.from_string(HAIRLINE)
        gridline.width = Pt(0.75)
        chart.category_axis.has_major_gridlines = False
    palette = brand.get("chart_series") or SERIES_COLOURS
    plot = chart.plots[0]
    if chart_spec["type"] == "pie":
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
    else:
        for i, series in enumerate(plot.series):
            colour = _rgb(palette[i % len(palette)])
            if chart_spec["type"] == "line":
                series.format.line.color.rgb = colour
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = colour
    if spec.get("takeaway"):
        _rect(slide, (9.3, 1.4, 3.4, 5.3), colours["light"], MSO_SHAPE.ROUNDED_RECTANGLE)
        (text,), pt = _fit([spec["takeaway"]], (3.0, 4.9), 18, warnings, [f"slide {number} takeaway"])
        _text(slide, (9.5, 1.6, 3.0, 4.9), text, pt, colours["ink"], brand)


RENDER = {"title": _title_slide, "messages": _messages_slide, "text_table": _text_table_slide,
          "metrics": _metrics_slide, "chart": _chart_slide}


def _save(prs, out_path):
    save_via_temp(out_path, prs.save, suffix=".pptx")


def build_deck(spec, out_path, brand=None):
    brand = brand or load_brand()
    spec, removed_chars = xmlsafe.clean_tree(spec)
    slides = spec.get("slides") if isinstance(spec, dict) else None
    if not isinstance(slides, list) or not slides:
        raise ToolkitError(("spec", "needs a non-empty 'slides' list"))
    if len(slides) > MAX_SLIDES:
        raise ToolkitError(("slides", f"a deck can have at most {MAX_SLIDES} slides; this one has {len(slides)}"))
    problems = []
    if "footer" in spec and spec["footer"] is not None and not isinstance(spec["footer"], str):
        problems.append(("footer", "footer must be a string"))
    problems += [p for i, s in enumerate(slides) for p in validate_slide(i, s)]
    if problems:
        raise ToolkitError(problems)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    warnings = []
    if removed_chars:
        warnings.append(("deck", f"removed {removed_chars} character(s) that PowerPoint cannot store"))
    for number, slide_spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        headline = None if slide_spec["layout"] == "title" else slide_spec["headline"]
        footer = slide_spec.get("footer") or spec.get("footer") or brand["footer_text"]
        _chrome(slide, number, headline, footer, brand, warnings)
        RENDER[slide_spec["layout"]](slide, slide_spec, number, brand, warnings)
    _save(prs, out_path)
    return {"deck": os.path.abspath(out_path), "slides": len(slides), "_warnings": warnings}
