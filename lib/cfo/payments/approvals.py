"""T7: the approval matrix -- the other thing the fraud checks compare an
invoice against. A row states a role, what that role may approve (free
text, kept for a person to read but not otherwise used here), a limit
amount and the currency it applies in. Real matrices arrive as a CSV, an
Excel sheet or a Word table (board papers are usually the latter two), so
`load_matrix` reads all three and gives back the same limits regardless of
which one a customer happened to send.

`load_matrix` is lenient about column headers -- "Limit", "Amount" and
"Approval limit" all mean the same thing to a person filling in a
spreadsheet by hand, and rejecting a real matrix over a header name this
tool did not anticipate would be a worse failure than reading it. See
`_HEADER_SYNONYMS`. It is not lenient about the values: a role with no
limit, or a limit that is not a number, is a `ToolkitError` -- this is the
reference data the fraud checks trust, not something to guess at.

**A run with no approval matrix must still work.** `load_matrix(None)`
returns `{}`, and so does a file that opens cleanly but states no rows (a
CSV with only a header, a sheet with none, a Word document with no table).
Neither is an error: there may genuinely be no matrix yet. What must never
happen silently is a limit-based check running against that empty matrix
and reporting "no problems found" as if it had actually checked something.
`available(matrix)` exists so a caller reads that as a fact -- `False` for
an absent or empty matrix -- rather than inferring it from an empty dict,
which looks exactly like "checked, and every limit was fine". T8's two
limit-based checks call `available` first and report themselves as **not
performed, by name**, when it is `False` -- the same discipline
`cfo.payments.validate.check_against_history` applies to a missing run
history, for the same reason: a check that quietly does nothing because it
had no data is worse than no check at all.

Amounts are `Decimal`, as everywhere else in this tool. A limit typed into
a spreadsheet as a number is unavoidably a `float` or an `int` by the time
openpyxl hands it back -- there is no more precise representation at that
source for this tool to insist on instead -- so, unlike
`cfo.payments.model` and `cfo.payments.validate`, a numeric cell value here
is accepted and converted with `Decimal(str(value))`. That is a different
situation from a caller's own Python code passing a literal `float`, which
had a choice and should have made a `Decimal` instead; nothing here relaxes
that rule for a caller.
"""
import csv
import os
import re
from decimal import Decimal, InvalidOperation

from cfo.console import ToolkitError

_CURRENCY_RE = re.compile(r"^[A-Z]{3}$")

# Real matrices call these columns all sorts of things; every key here is
# matched case-insensitively, with runs of whitespace collapsed, against a
# header cell. An unrecognised header is not an error -- the sheet may carry
# other columns (a date, a sign-off name) this tool has no use for -- it is
# simply not one of the four fields a row needs.
_HEADER_SYNONYMS = {
    "role": "role",
    "approver": "role",
    "approving role": "role",
    "approver role": "role",
    "may_approve": "scope",
    "may approve": "scope",
    "what they may approve": "scope",
    "approves": "scope",
    "scope": "scope",
    "description": "scope",
    "limit": "limit",
    "amount": "limit",
    "limit amount": "limit",
    "approval limit": "limit",
    "max amount": "limit",
    "currency": "currency",
    "ccy": "currency",
}


def _normalise_header(header):
    return re.sub(r"\s+", " ", str(header or "").strip().casefold())


def _canonicalise(raw_row):
    row = {}
    for header, value in raw_row.items():
        canonical = _HEADER_SYNONYMS.get(_normalise_header(header))
        if canonical:
            row[canonical] = value
    return row


def _parse_amount(value, where):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, (where, "limit amount is required")
    if isinstance(value, bool):
        return None, (where, f"limit amount {value!r} is not a number")
    if isinstance(value, Decimal):
        return value, None
    if isinstance(value, (int, float)):
        # See the module docstring: a spreadsheet cell has no more precise
        # representation than this to offer.
        return Decimal(str(value)), None
    try:
        return Decimal(str(value).strip()), None
    except InvalidOperation:
        return None, (where, f"limit amount {value!r} is not a number")


def _row_to_limit(raw_row, where):
    row = _canonicalise(raw_row)
    role = str(row.get("role") or "").strip()
    scope = str(row.get("scope") or "").strip()
    currency = str(row.get("currency") or "").strip().upper()
    if not role:
        raise ToolkitError((where, "needs a role"))
    if not _CURRENCY_RE.match(currency):
        raise ToolkitError((where, f"currency {currency!r} is not an ISO 4217 three-letter "
                                   "code"))
    limit, err = _parse_amount(row.get("limit"), where)
    if err:
        raise ToolkitError(err)
    return role, currency, limit, scope


def _read_csv_rows(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            return []
        return [dict(row) for row in reader if any(str(v or "").strip() for v in row.values())]


def _read_xlsx_rows(path):
    from openpyxl import load_workbook  # lazy: a CSV-only caller needs no openpyxl at all
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(h or "").strip() for h in next(rows)]
        except StopIteration:
            return []
        result = []
        for values in rows:
            if not any(v is not None and str(v).strip() for v in values):
                continue
            result.append(dict(zip(header, values)))
        return result
    finally:
        workbook.close()


def _read_docx_rows(path):
    from docx import Document  # lazy: a CSV-only caller needs no python-docx at all
    document = Document(path)
    if not document.tables:
        return []
    table = document.tables[0]
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return []
    header, data_rows = rows[0], rows[1:]
    return [dict(zip(header, values)) for values in data_rows if any(v for v in values)]


def load_matrix(path):
    """`{role: [{"currency", "limit", "scope"}, ...]}` read from `path`
    (CSV, Excel or Word), or `{}` when `path` is `None` or the file states
    no rows -- see the module docstring for why neither is an error, and
    why `available` exists so a caller never has to tell those two cases
    apart from the empty dict alone."""
    if path is None:
        return {}
    if not os.path.isfile(path):
        raise ToolkitError(("path", f"approval matrix file not found: {path}"))
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        rows = _read_csv_rows(path)
    elif ext in (".xlsx", ".xlsm"):
        rows = _read_xlsx_rows(path)
    elif ext == ".docx":
        rows = _read_docx_rows(path)
    else:
        raise ToolkitError(("path", f"unsupported approval matrix file type '{ext}'; "
                                    "use .csv, .xlsx or .docx"))
    matrix = {}
    for i, raw_row in enumerate(rows, start=2):  # row 1 is the header
        role, currency, limit, scope = _row_to_limit(raw_row, f"row {i}")
        matrix.setdefault(role, []).append({"currency": currency, "limit": limit, "scope": scope})
    return matrix


def _find_role(role, matrix):
    target = str(role or "").strip().casefold()
    for key in matrix:
        if key.strip().casefold() == target:
            return key
    return None


def limit_for(role, matrix, currency):
    """The `Decimal` limit `role` may approve in `currency`, or `None` when
    the role is not in `matrix` or states no limit in that currency. Role
    and currency are matched case- and whitespace-insensitively -- this is
    a control lookup, not the supplier-name matching in
    `cfo.payments.suppliers`, so there is no "near-exact only" concern here:
    a role is either configured or it is not."""
    key = _find_role(role, matrix)
    if key is None:
        return None
    currency = str(currency or "").strip().upper()
    for entry in matrix[key]:
        if entry["currency"] == currency:
            return entry["limit"]
    return None


def limits(matrix):
    """Every `(role, currency, amount)` in `matrix`, sorted by role then
    currency."""
    rows = [(role, entry["currency"], entry["limit"])
            for role, entries in matrix.items() for entry in entries]
    return sorted(rows, key=lambda row: (row[0].casefold(), row[1]))


def available(matrix):
    """`False` for an absent or empty matrix, `True` otherwise -- see the
    module docstring for why a caller must read this rather than infer it
    from `matrix` being falsy itself."""
    return bool(matrix)
